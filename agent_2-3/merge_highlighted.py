#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将所有患者的 _highlighted.xlsx 合并成一个汇总的高亮 Excel 文件
"""
import os
import sys
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
import tqdm

def merge_highlighted_excels(filled_tables_dir: str, output_path: str):
    """
    合并所有患者的 _highlighted.xlsx 到一个文件

    Args:
        filled_tables_dir: filled_tables 目录路径
        output_path: 输出的汇总 Excel 路径
    """
    # 样式定义
    YELLOW = PatternFill("solid", fgColor="FFFF00")
    GRAY   = PatternFill("solid", fgColor="F2F2F2")
    WHITE  = PatternFill("solid", fgColor="FFFFFF")
    BOLD9  = Font(bold=True, size=9)
    NORM9  = Font(size=9)
    WRAP   = Alignment(wrap_text=True)

    # 查找所有 _highlighted.xlsx 文件
    highlighted_files = []
    for root, dirs, files in os.walk(filled_tables_dir):
        for f in files:
            if f.endswith("_highlighted.xlsx"):
                highlighted_files.append(os.path.join(root, f))

    if not highlighted_files:
        print("未找到任何 _highlighted.xlsx 文件")
        return

    print(f"找到 {len(highlighted_files)} 个患者的高亮文件")

    # 创建新的工作簿
    merged_wb = Workbook()
    merged_ws = merged_wb.active
    merged_ws.title = "填写结果汇总"

    # 用于存储列宽
    col_widths = {}
    header_written = False
    current_row = 1

    # 逐个读取并合并
    for file_path in tqdm.tqdm(highlighted_files, desc="合并进度", unit="人"):
        try:
            wb = load_workbook(file_path, data_only=False)
            ws = wb["填写结果"]

            # 第一个文件：写入表头
            if not header_written:
                for col_idx in range(1, ws.max_column + 1):
                    header_cell = ws.cell(row=1, column=col_idx)
                    new_cell = merged_ws.cell(row=1, column=col_idx, value=header_cell.value)
                    new_cell.font = BOLD9
                    new_cell.alignment = WRAP
                    # 记录列宽
                    col_letter = get_column_letter(col_idx)
                    if col_letter in ws.column_dimensions:
                        col_widths[col_idx] = ws.column_dimensions[col_letter].width
                header_written = True
                current_row = 2

            # 复制数据行（第2行）及其格式
            for col_idx in range(1, ws.max_column + 1):
                src_cell = ws.cell(row=2, column=col_idx)
                dst_cell = merged_ws.cell(row=current_row, column=col_idx, value=src_cell.value)

                # 复制样式
                if src_cell.fill and src_cell.fill.fgColor:
                    fg_color = src_cell.fill.fgColor.rgb
                    if fg_color:
                        # 黄色背景（新填入）
                        if "FFFF00" in str(fg_color):
                            dst_cell.fill = YELLOW
                            dst_cell.font = BOLD9
                        # 灰色背景（空值）
                        elif "F2F2F2" in str(fg_color):
                            dst_cell.fill = GRAY
                            dst_cell.font = NORM9
                        # 白色背景（原有值）
                        else:
                            dst_cell.fill = WHITE
                            dst_cell.font = NORM9
                    else:
                        dst_cell.font = NORM9
                else:
                    dst_cell.font = NORM9

                dst_cell.alignment = WRAP

            current_row += 1
            wb.close()

        except Exception as e:
            print(f"处理文件失败 {file_path}: {e}")
            continue

    # 设置列宽
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        merged_ws.column_dimensions[col_letter].width = width

    # 保存
    merged_wb.save(output_path)
    print(f"\n汇总完成！共 {current_row - 2} 个患者")
    print(f"输出文件: {output_path}")
    print(f"文件大小: {os.path.getsize(output_path) / 1024 / 1024:.2f} MB")


if __name__ == "__main__":
    filled_tables_dir = "/mnt/GPU_10T/home/xkc/my_BioDSA/dataprocess/agent_2-3/results-all/results_20260512_001514/filled_tables"
    output_path = os.path.join(filled_tables_dir, "filled_merged_highlighted.xlsx")

    merge_highlighted_excels(filled_tables_dir, output_path)
