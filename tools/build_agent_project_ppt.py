from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


TEMPLATE = Path(
    "/Users/mkbk/Library/Containers/com.tencent.xinWeChat/Data/Documents/"
    "xwechat_files/wxid_u85c5mjfdipn12_b063/temp/RWTemp/2026-05/"
    "d69a6d81-d8dd-4ff3-af59-2f0202f8a18a/科研 总结-陈俊华.pptx"
)
OUTPUT = Path("/Users/mkbk/PycharmProjects/new/智能体项目汇报-3页版.pptx")

FONT = "Microsoft YaHei"
DARK = RGBColor(34, 53, 91)
ACCENT = RGBColor(67, 90, 146)
LIGHT = RGBColor(239, 244, 251)
WHITE = RGBColor(255, 255, 255)


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def fill_box(shape, lines: list[str], font_size: int = 14, bold_first: bool = False) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.LEFT
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(font_size)
            run.font.color.rgb = DARK
            run.font.bold = bold_first and idx == 0


def set_shape_text(shape, text: str, font_size: int = 12, bold: bool = True) -> None:
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = text
    for run in p.runs:
        run.font.name = FONT
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = DARK


def style_new_shape(shape, fill_rgb: RGBColor, font_rgb: RGBColor = DARK) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.color.rgb = ACCENT
    tf = shape.text_frame
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.name = FONT
            run.font.color.rgb = font_rgb


def set_table_cell(cell, text: str, size: int = 11, bold: bool = False, center: bool = True) -> None:
    tf = cell.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    for run in p.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = DARK


def build_slide1(slide) -> None:
    slide.shapes[4].text = "医疗智能体项目：为什么要做这条可追踪、可验证的数据处理链路"

    fill_box(
        slide.shapes[18],
        [
            "一、背景",
            "• 医疗多模态原始数据同时包含文档、表格、图像和半结构化字段，人工整理成本高且容易出错。",
            "• 下游建模是否可靠，取决于上游重组、抽取、裁剪、清洗是否稳定，而这一步通常最难复现。",
            "• 真实项目里最常见的问题不是“模型不会学”，而是 records 不完整、模态识别不稳、配置路径易错、输出口径不一致。",
        ],
        font_size=14,
        bold_first=True,
    )
    fill_box(
        slide.shapes[5],
        [
            "二、核心问题",
            "能否把医疗原始目录整理、结构化抽取、任务裁剪、数据修复与训练集生成串成一个可恢复、可验证、可审计的智能体流水线？",
        ],
        font_size=16,
        bold_first=True,
    )
    fill_box(
        slide.shapes[19],
        [
            "三、研究意义",
            "• 工程意义：把“人肉清数据”变成带状态、带验证器、可续跑的自动化流程。",
            "• 方法意义：把 Agent 推理、工具调用、规则校验和记忆机制组合到同一条链路中。",
            "• 应用意义：提升真实医疗数据进入分析与建模阶段的效率，减少人工整理成本与口径偏差。",
            "• 应用意义：为死亡预测、离院去向预测等下游临床任务提供更稳定、可追溯的数据输入，增强结果可信度。",
        ],
        font_size=14,
        bold_first=True,
    )

    pipeline_boxes = {
        "top": slide.shapes[23],
        "upper_mid": slide.shapes[22],
        "lower_mid": slide.shapes[21],
        "bottom": slide.shapes[20],
    }
    removable = [slide.shapes[idx] for idx in range(17, 5, -1)]
    for shape in removable:
        remove_shape(shape)

    pipeline_labels = {
        "top": "原始病历 / 图像 / 表格",
        "upper_mid": "Step1 重组与模态标注\n生成 records.json",
        "lower_mid": "Step2-5 抽取、裁剪、清洗\n得到可用特征表",
        "bottom": "Step6-7 一致性验证\n输出 ML 数据集",
    }
    for key, text in pipeline_labels.items():
        shape = pipeline_boxes[key]
        set_shape_text(shape, text, font_size=12)
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT
        shape.line.color.rgb = ACCENT

    box1 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(10.75), Inches(2.88), Inches(0.95), Inches(0.42))
    style_new_shape(box1, RGBColor(222, 232, 248))
    set_shape_text(box1, "规则记忆", font_size=10)

    box2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(10.75), Inches(3.40), Inches(0.95), Inches(0.42))
    style_new_shape(box2, RGBColor(222, 232, 248))
    set_shape_text(box2, "案例记忆", font_size=10)


def build_slide2(slide) -> None:
    slide.shapes[4].text = "项目模型：多阶段 Agent + 校验器 + Playbook/ReMe 记忆"

    fill_box(
        slide.shapes[7],
        ["一、总体设计与记忆机制"],
        font_size=16,
        bold_first=True,
    )
    fill_box(
        slide.shapes[5],
        [
            "• 状态机编排：支持 full_pipeline、step1_only、resume_from_records 等多入口。",
            "• Agent 与工具解耦：领域判断交给 Agent，数据读写与保存交给确定性工具。",
            "• 每一步后接 Validator：不只追求“跑通”，更强调 records、CSV、JSON 产物是否满足契约。",
        ],
        font_size=14,
    )
    fill_box(
        slide.shapes[6],
        [
            "• Playbook Memory：保存流程约束、失败模式、输出契约，适合规则型经验。",
            "• ReMe Case Memory：保存 Step1 脚本案例、输入 profile、执行摘要，适合相似案例复用。",
            "• 记忆在任务前检索、任务后反思更新，形成“运行 - 反馈 - 归档”闭环。",
            "• 关键链路：Step1 → Step2/3 → Step4 → Step5 → Step6/7。",
        ],
        font_size=14,
    )

    label_map = {
        8: "原始目录 / records.json",
        9: "Step1\n重组与模态分析",
        10: "下游 ML / 分析",
        14: "Playbook\n规则记忆",
        19: "Step2/3\n结构化抽取",
        20: "Step4\n任务裁剪",
        21: "ReMe\n案例记忆",
        22: "Step5\n数据修复",
        23: "Step6/7\n验证与数据集生成",
        25: "记忆注入",
        26: "可追踪流水线",
        27: "Validator\n失败回环",
    }
    for idx, text in label_map.items():
        set_shape_text(slide.shapes[idx], text, font_size=10 if idx in {14, 22, 26, 28} else 11)


def build_slide3(slide) -> None:
    slide.shapes[4].text = "当前结果：已经形成可验收产物，但端到端稳定性仍需继续打磨"

    fill_box(slide.shapes[6], ["一、当前结果"], font_size=16, bold_first=True)
    fill_box(
        slide.shapes[5],
        [
            "二、当前实验结论",
            "",
            "发现 1：项目已经具备分阶段输出与可验证接口",
            "当前仓库内可直接看到 records、记忆测试、Step7 输出契约三类证据，说明链路不是纯概念设计。",
            "",
            "发现 2：重点价值不只是自动化，更是可追踪与可恢复",
            "同一条链路支持 records 续跑、模块单独验证、记忆检索辅助与输出格式统一，这对真实医疗数据工程更关键。",
        ],
        font_size=13,
        bold_first=True,
    )
    fill_box(
        slide.shapes[8],
        [
            "三、Discussion",
            "",
            "当前优势",
            "• 任务入口清晰，分步产物可检查。",
            "• 规则记忆与案例记忆分层，便于复用经验而不污染硬约束。",
            "",
            "当前问题",
            "• 真正的大规模端到端结果仍依赖真实数据与配置环境。",
            "• Step7 面向 outcome task 时仍需重点防 label leakage。",
            "• YAML / 路径 / 本地依赖仍可能影响稳定复现。",
        ],
        font_size=13,
        bold_first=True,
    )
    fill_box(
        slide.shapes[9],
        [
            "下一步工作：",
            "1、在真实数据上做完整 benchmark，补齐每一步耗时、通过率和失败类型统计。",
            "2、把 ReMe 记忆继续向 Step1 脚本复用与异常修复建议扩展。",
            "3、补强 Step7 outcome 任务的泄漏检查，形成更严格的数据验收门槛。",
            "4、继续收敛本地配置与路径依赖，减少“环境一变就跑不通”的问题。",
        ],
        font_size=13,
        bold_first=True,
    )

    table = slide.shapes[7].table
    rows = [
        ["模块", "验证证据", "当前观测", "说明", "主要风险", "结论"],
        ["Step1", "records.json", "34 条记录", "当前仓库样例中已存在文件级记录", "真实全量数据仍需复跑", "已形成基础重组产物"],
        ["ReMe", "pytest", "5/5 通过", "记忆工具单测通过，可独立验证", "尚未完全代表主链路表现", "独立模块可用"],
        ["Step7", "源码契约", "CSV/JSONL/JSON", "输出包含 label_name、y、label_ICD10", "需继续防 outcome leakage", "下游训练接口已标准化"],
    ]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            set_table_cell(table.cell(r, c), value, size=10 if r else 11, bold=(r == 0))


def main() -> None:
    prs = Presentation(str(TEMPLATE))
    build_slide1(prs.slides[0])
    build_slide2(prs.slides[1])
    build_slide3(prs.slides[2])
    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
